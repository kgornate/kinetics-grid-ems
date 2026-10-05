#!/usr/bin/env python3
"""Build a Kinetics-compatible protocol catalog from Lineage BMS2EMS V05.

The workbook's Decimal address column is treated as authoritative because some
Hexadecimal cells are formulas or were imported by Excel as scientific notation.
No attempt is made to invent missing protocol semantics: unknown Unit-ID assignment
for auxiliary devices and FLOAT word order are recorded as commissioning gaps.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SHEETS: dict[str, dict[str, Any]] = {
    "System Data": {"scope": "bank", "category": "system_data", "read_function": 4, "poll": "normal"},
    "System Alarm": {"scope": "bank", "category": "system_alarm", "read_function": 2, "poll": "fast"},
    "Rack Data": {"scope": "rack", "category": "rack_data", "read_function": 4, "poll": "normal"},
    "Rack Alarm": {"scope": "rack", "category": "rack_alarm", "read_function": 2, "poll": "fast"},
    "Cell Temp": {"scope": "rack", "category": "cell_temperature", "read_function": 4, "poll": "bulk"},
    "Cell Balance": {"scope": "rack", "category": "cell_balance", "read_function": 2, "poll": "bulk"},
    "Cell Volt": {"scope": "rack", "category": "cell_voltage", "read_function": 4, "poll": "bulk"},
    "Cell SOC": {"scope": "rack", "category": "cell_soc", "read_function": 4, "poll": "bulk"},
    "Pack Temp": {"scope": "rack", "category": "pack_temperature", "read_function": 4, "poll": "bulk"},
    "Pack Volt": {"scope": "rack", "category": "pack_voltage", "read_function": 4, "poll": "bulk"},
    "Temp&Humidity": {"scope": "environment", "category": "temp_humidity", "read_function": 4, "poll": "normal"},
    "Water sensor": {"scope": "environment", "category": "water_sensor", "read_function": 4, "poll": "normal"},
    "Chiller": {"scope": "environment", "category": "chiller", "read_function": 4, "poll": "normal"},
    "Dehumidification": {"scope": "environment", "category": "dehumidifier", "read_function": 4, "poll": "normal"},
    "Setting": {"scope": "bank", "category": "setting", "read_function": 3, "poll": "disabled"},
}

SYSTEM_KEY_OVERRIDES = {
    0: "protocol_version",
    1: "battery_system_state",
    2: "installed_rack_count",
    3: "operating_rack_count",
    4: "operating_rack_bitmap",
    5: "max_soc_rack_id",
    6: "min_soc_rack_id",
    7: "max_voltage_rack_id",
    8: "min_voltage_rack_id",
    9: "max_cell_voltage_rack_id",
    10: "max_cell_voltage_cell_id",
    11: "min_cell_voltage_rack_id",
    12: "min_cell_voltage_cell_id",
    13: "max_cell_temperature_rack_id",
    14: "max_cell_temperature_cell_id",
    15: "min_cell_temperature_rack_id",
    16: "min_cell_temperature_cell_id",
    19: "emergency_stop_input",
    20: "busbar_cabinet_door_input",
    21: "aux_220v_switch_feedback",
    22: "aux_220v_switch_alarm",
    23: "container_door_input",
    24: "ups_fault_input",
    25: "fss_warning_input",
    26: "fss_alarm_input",
    27: "fss_fault_input",
    28: "fan_fault_input",
    29: "gas_detector_fault_input",
    30: "aerosol_spray_feedback",
    51: "system_voltage",
    53: "system_current",
    55: "system_soc",
    57: "system_soh",
    59: "system_charge_energy_session",
    61: "system_discharge_energy_session",
    63: "system_full_charge_capacity",
    65: "system_full_discharge_capacity",
    67: "system_accumulated_charge_energy",
    69: "system_accumulated_discharge_energy",
    71: "system_cell_voltage_delta",
    73: "system_max_cell_voltage",
    75: "system_min_cell_voltage",
    77: "system_cell_temperature_delta",
    79: "system_max_cell_temperature",
    81: "system_min_cell_temperature",
    83: "system_rack_soc_delta",
    85: "system_max_rack_soc",
    87: "system_min_rack_soc",
    89: "system_rack_voltage_delta",
    91: "system_max_rack_voltage",
    93: "system_min_rack_voltage",
    95: "system_max_charge_power",
    97: "system_max_discharge_power",
    99: "system_max_charge_current",
    101: "system_max_discharge_current",
}

RACK_KEY_OVERRIDES = {
    1: "rack_state",
    2: "rack_operating_state",
    3: "rack_cell_count",
    4: "rack_temperature_sensor_count",
    5: "rack_pack_count",
    6: "rack_charge_count",
    7: "rack_discharge_count",
    8: "rack_max_cell_voltage_cell_id",
    9: "rack_min_cell_voltage_cell_id",
    10: "rack_max_cell_temperature_cell_id",
    11: "rack_min_cell_temperature_cell_id",
    12: "rack_max_cell_internal_resistance_cell_id",
    13: "rack_min_cell_internal_resistance_cell_id",
    14: "rack_max_cell_soc_cell_id",
    15: "rack_min_cell_soc_cell_id",
    16: "rack_max_cell_soh_cell_id",
    17: "rack_min_cell_soh_cell_id",
    18: "positive_contactor_feedback",
    19: "negative_contactor_feedback",
    20: "mcb_feedback",
    21: "fuse_feedback",
    51: "rack_voltage",
    53: "rack_current",
    55: "rack_soc",
    57: "rack_soh",
    59: "hv_box_temperature",
    61: "positive_insulation_resistance",
    63: "negative_insulation_resistance",
    65: "rack_charge_energy_session",
    67: "rack_discharge_energy_session",
    69: "rack_full_charge_capacity",
    71: "rack_full_discharge_capacity",
    73: "rack_accumulated_charge_energy",
    75: "rack_accumulated_discharge_energy",
    77: "rack_cell_voltage_delta",
    79: "rack_average_cell_voltage",
    81: "rack_max_cell_voltage",
    83: "rack_min_cell_voltage",
    85: "rack_cell_temperature_delta",
    87: "rack_average_cell_temperature",
    89: "rack_max_cell_temperature",
    91: "rack_min_cell_temperature",
    93: "rack_average_cell_internal_resistance",
    95: "rack_max_cell_internal_resistance",
    97: "rack_min_cell_internal_resistance",
    99: "rack_average_cell_soc",
    101: "rack_max_cell_soc",
    103: "rack_min_cell_soc",
    105: "rack_average_cell_soh",
    107: "rack_max_cell_soh",
    109: "rack_min_cell_soh",
    111: "rack_max_charge_power",
    113: "rack_max_discharge_power",
    115: "rack_max_charge_current",
    117: "rack_max_discharge_current",
    119: "hv_box_temperature_2",
    121: "hv_box_temperature_3",
    123: "hv_box_temperature_4",
}

SYSTEM_ALARM_KEYS = {
    1: "emergency_stop",
    3: "fire_start",
    4: "fire_fault",
    5: "water_alarm",
    6: "combustible_gas_alarm",
    7: "container_door_alarm",
    8: "pcs_communication_fault",
    9: "ems_communication_fault",
    10: "dc_combiner_busbar_switch_state",
}

SETTING_BANK_KEYS = {
    1: "dc_combiner_open_delay_seconds",
    2: "rtc_year",
    3: "rtc_month",
    4: "rtc_day",
    5: "rtc_hour",
    6: "rtc_minute",
    7: "rtc_second",
    8: "rtc_set_confirmation",
    11: "fault_reset",
}

SETTING_RACK_KEYS = {
    1: "positive_contactor_command",
    2: "negative_contactor_command",
}

STATE_ENUM = {
    "4369": "NoChg",       # 0x1111
    "8738": "NoDchg",     # 0x2222
    "21845": "Standby",   # 0x5555
    "43690": "Fault",     # 0xAAAA
    "48059": "Normal",    # 0xBBBB
    "52428": "Warn",      # 0xCCCC
}


def slug(text: str) -> str:
    text = text.replace("℃", " temperature ").replace("%", " percent ")
    text = text.replace("&", " and ")
    text = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower()
    return re.sub(r"_+", "_", text) or "unnamed"


def normalize_dtype(raw: Any) -> tuple[str, str | None]:
    source = str(raw or "").strip().upper()
    if source == "FLOAT":
        return "FLOAT32", None
    if source == "UINT16":
        return "U16", None
    if source == "INT16":
        return "S16", None
    if source == "UINT32":
        return "U32", None
    # The workbook contains UINT17..UINT33 in a chiller write section. These
    # are not valid Modbus scalar types and are treated as a documentation typo.
    if re.fullmatch(r"UINT(?:1[7-9]|2[0-9]|3[0-3])", source):
        return "U16", f"source workbook data type {source} treated as UINT16; vendor confirmation required"
    if not source:
        return "U16", "source workbook data type missing"
    return "U16", f"unrecognized source workbook data type {source}; treated as UINT16"



def point_key(sheet: str, address: int, meaning: str, scope: str, *, setting_scope: str | None = None) -> str:
    if sheet == "System Data" and address in SYSTEM_KEY_OVERRIDES:
        return SYSTEM_KEY_OVERRIDES[address]
    if sheet == "Rack Data" and address in RACK_KEY_OVERRIDES:
        return RACK_KEY_OVERRIDES[address]
    if sheet == "System Alarm" and address in SYSTEM_ALARM_KEYS:
        return SYSTEM_ALARM_KEYS[address]
    if sheet == "Setting":
        if setting_scope == "rack" and address in SETTING_RACK_KEYS:
            return SETTING_RACK_KEYS[address]
        if setting_scope == "bank" and address in SETTING_BANK_KEYS:
            return SETTING_BANK_KEYS[address]
    patterns = [
        (r"^(\d+)#\s*cell\s+voltage$", "cell_voltage_{}"),
        (r"^(\d+)#\s*cell\s+temperature$", "cell_temperature_{}"),
        (r"^(\d+)#\s*cell\s+soc$", "cell_soc_{}"),
        (r"^(\d+)#\s*cell\s+charge\s+balance\s+state$", "cell_charge_balance_{}"),
        (r"^(\d+)#\s*cell\s+discharge\s+balance\s+state$", "cell_discharge_balance_{}"),
        (r"^(\d+)#\s*pack\s+voltage$", "pack_voltage_{}"),
        (r"^(\d+)#\s*pack\s+negative\s+pole\s+temperature$", "pack_negative_pole_temperature_{}"),
        (r"^(\d+)#\s*pack\s+positive\s+pole\s+temperature$", "pack_positive_pole_temperature_{}"),
    ]
    clean = re.sub(r"\s+", " ", meaning.strip(), flags=re.MULTILINE)
    for pattern, template in patterns:
        match = re.match(pattern, clean, flags=re.I)
        if match:
            return template.format(match.group(1))
    return f"{slug(meaning)}_{address}" if meaning.lower() in {"fault code", "reserve", "reserved"} else slug(meaning)


def poll_class(sheet: str, address: int, meaning: str, reserved: bool) -> str:
    if reserved:
        return "disabled"
    if sheet in {"System Alarm", "Rack Alarm"}:
        return "fast"
    if sheet == "System Data" and address in {1, 3, 4, 19, 51, 53, 55, 95, 97, 99, 101}:
        return "fast"
    if sheet == "Rack Data" and address in {1, 2, 18, 19, 20, 21, 51, 53, 55, 61, 63, 111, 113, 115, 117}:
        return "fast"
    return SHEETS[sheet]["poll"]


def enum_for(sheet: str, address: int, meaning: str, remarks: str) -> dict[str, str]:
    if (sheet == "System Data" and address == 1) or (sheet == "Rack Data" and address == 1):
        return dict(STATE_ENUM)
    if sheet == "Rack Data" and address == 2:
        return {"1": "Open", "2": "Standby", "3": "Charge", "4": "Discharge"}
    if sheet == "Rack Data" and address in {18, 19, 20, 21, 22, 23}:
        return {"0": "Open", "1": "Closed"}
    if sheet == "Setting" and meaning and "contactor" in meaning.lower():
        return {"0": "BMS automatic control", "1": "EMS forced close", "2": "EMS forced open"}
    if "0--normal" in remarks.lower() and "1--" in remarks.lower():
        return {"0": "Normal", "1": "Active"}
    return {}


def build(workbook: Path) -> dict[str, Any]:
    wb = load_workbook(workbook, data_only=False, read_only=True)
    points: list[dict[str, Any]] = []
    reserved_ranges: list[dict[str, Any]] = []

    for sheet, spec in SHEETS.items():
        ws = wb[sheet]
        setting_scope = "bank"
        for row_index, row_tuple in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            row = list(row_tuple)
            register_type = str(row[0] or "")
            if sheet == "Setting" and "Modbus address 2" in register_type:
                setting_scope = "rack"

            address_raw = row[2] if len(row) > 2 else None
            if not isinstance(address_raw, (int, float)):
                continue
            address = int(address_raw)
            access_raw = str(row[3] or "RO").upper().strip()
            meaning = str(row[4] or "").strip()
            source_dtype = row[5] if len(row) > 5 else None
            coefficient = row[6] if len(row) > 6 else None
            offset = row[7] if len(row) > 7 else None
            unit = None
            # Most sheets put units at column I; Cell Temp/Balance variants use
            # slightly different columns but column I remains the primary source.
            if len(row) > 8 and isinstance(row[8], str):
                unit = row[8].strip() or None
            remarks_candidates = [str(value) for value in row[9:] if value not in (None, "")]
            remarks = " | ".join(remarks_candidates)
            reserved = (not meaning) or ("reserve" in meaning.lower())

            scope = spec["scope"]
            category = spec["category"]
            read_function = int(spec["read_function"])
            if sheet == "Setting":
                scope = setting_scope
                read_function = 3
            if sheet == "Chiller" and ("Write Chiller data" in register_type or access_raw in {"RW", "W"}):
                read_function = 3

            dtype, dtype_note = normalize_dtype(source_dtype)
            width = 2 if dtype in {"FLOAT32", "U32", "S32"} else 1
            key = point_key(sheet, address, meaning or "reserved", scope, setting_scope=setting_scope)
            point: dict[str, Any] = {
                "id": f"{scope}.{category}.{key}",
                "sheet": sheet,
                "scope": scope,
                "category": category,
                "address": address,
                "address_hex": f"0x{address:04X}",
                "key": key,
                "name_en": meaning or None,
                "name_cn": None,
                "access": "RW" if access_raw in {"RW", "WR"} else ("W" if access_raw == "W" else "R"),
                "data_type": dtype,
                "register_width": width,
                "element_count": 1,
                "register_count": width,
                "scale": float(coefficient) if isinstance(coefficient, (int, float)) else None,
                "offset": float(offset) if isinstance(offset, (int, float)) else None,
                "unit": unit,
                "range_text": remarks if "Value Range" in remarks else None,
                "enum": enum_for(sheet, address, meaning, remarks),
                "read_function": read_function,
                "write_function": 6 if "W" in access_raw else None,
                "write_functions": [6] if "W" in access_raw else [],
                "poll_class": poll_class(sheet, address, meaning, reserved),
                "poll_enabled": not reserved and poll_class(sheet, address, meaning, reserved) != "disabled",
                "reserved": reserved,
                "source": "Lineage-BMS2EMS-Protocol-1 V05_final.xlsx",
                "source_row": row_index,
                "source_data_type": str(source_dtype or ""),
                "notes": remarks or None,
            }
            if dtype_note:
                point["documentation_issue"] = dtype_note
            if sheet == "Cell Temp" and len(row) > 9:
                point["physical_mapping"] = row[9] or None
                if len(row) > 10:
                    point["display_mapping"] = row[10] or None
            if category in {"temp_humidity", "water_sensor", "chiller", "dehumidifier"}:
                point["endpoint_group"] = category
                point["unit_id_note"] = "Workbook says Modbus address starting from 101; exact auxiliary Unit-ID allocation requires commissioning/vendor confirmation."
            if read_function in {1, 2}:
                point["register_width"] = 1
                point["register_count"] = 1
                point["data_type"] = "U16"
                point["scale"] = None
                point["offset"] = None
            points.append(point)

    # Ensure keys are unique within each scope. Preserve canonical keys; append
    # address only when the workbook repeats a generic label.
    seen: set[tuple[str, str]] = set()
    for point in points:
        marker = (point["scope"], point["key"])
        if marker in seen:
            point["key"] = f"{point['key']}_{point['address']}"
            point["id"] = f"{point['scope']}.{point['category']}.{point['key']}"
        seen.add((point["scope"], point["key"]))

    active = [point for point in points if not point["reserved"]]
    metadata = {
        "source_file": workbook.name,
        "vendor": "Lineage",
        "protocol_version": "V05",
        "protocols": ["modbus_tcp", "modbus_rtu"],
        "supported_modbus_functions": [2, 3, 4, 6],
        "byte_order": "big_per_modbus",
        "float_word_order": "vendor_not_specified_configurable",
        "canonical_address_source": "Decimal column",
        "system_unit_id": 1,
        "rack_unit_id_rule": "Rack 1 begins at Unit ID 2; subsequent racks increment by one per workbook note",
        "auxiliary_unit_id_note": "Auxiliary sheets state Modbus address starting from 101; exact per-device Unit-ID allocation is not explicit in V05.",
        "commissioning_gaps": [
            "FLOAT32 word order",
            "battery/rack current sign convention",
            "0-based vs library address offset confirmation",
            "auxiliary device Unit-ID allocation",
            "rack contactor write behavior and safe sequence",
            "precharge ownership/sequence",
        ],
        "point_count": len(points),
        "active_point_count": len(active),
    }
    counts: dict[str, int] = {}
    for point in active:
        group = f"{point['scope']}.{point['category']}"
        counts[group] = counts.get(group, 0) + 1
    return {"metadata": metadata, "counts": counts, "points": points, "reserved_ranges": reserved_ranges}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    payload = build(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Wrote {args.output}: {len(payload['points'])} points, {payload['metadata']['active_point_count']} active")


if __name__ == "__main__":
    main()
